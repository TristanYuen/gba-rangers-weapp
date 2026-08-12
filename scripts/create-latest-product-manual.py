from __future__ import annotations

from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import Color, HexColor, white
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "tmp" / "product-manual-2026-07"
SCREENS = WORK / "screens"
OPTIMIZED = WORK / "optimized"
OUTPUT = ROOT / "output" / "product-manual"
PDF_OUT = OUTPUT / "大湾区流浪者小程序产品手册-2026年7月.pdf"
DOCX_OUT = OUTPUT / "大湾区流浪者小程序产品手册-2026年7月.docx"
LOGO = ROOT / "src" / "assets" / "gba-crest-white.png"

PW, PH = landscape(A4)
NAVY = HexColor("#071A52")
MIDNIGHT = HexColor("#06133D")
ROYAL = HexColor("#1848C8")
BLUE = HexColor("#3477EB")
ICE = HexColor("#BFD6FF")
PAPER = HexColor("#F3F6FC")
INK = HexColor("#0A173A")
MUTED = HexColor("#5F6D89")
LINE = HexColor("#D7E1F1")
GREEN = HexColor("#168B57")
GOLD = HexColor("#9C6D18")
RED = HexColor("#B3364D")


FEATURE_INDEX = [
    ("01", "双入口导航", "比赛中心、我的球队"),
    ("02", "成员与权限", "邀请码、审核、四级身份"),
    ("03", "比赛中心", "报名赛事、历史记录"),
    ("04", "赛事详情", "地图、集合、球衣、备注"),
    ("05", "队内报名", "参加、待定、缺席、候补"),
    ("06", "报名管理", "审核、提醒、筛选、名单复制"),
    ("07", "球队通知", "模板、受众、送达统计"),
    ("08", "球员名册", "位置分组、现役与历史成员"),
    ("09", "球员档案", "数据、隐私、头像、荣誉"),
    ("10", "数据榜单", "出场、进球、助攻"),
    ("11", "球队年鉴", "比赛章节、相册、精选"),
    ("12", "会费管理", "台账、档次、缴费状态"),
    ("13", "会费审批", "管理员申请、队长确认"),
    ("14", "运营中心", "任务看板、操作日志"),
    ("15", "成员与媒体审核", "绑定档案、封面、精选"),
    ("16", "赛事运营", "创建、发布、赛后统计"),
]


SECTIONS = [
    {
        "eyebrow": "PRODUCT STRUCTURE",
        "title": "从比赛到球队运营的双入口",
        "subtitle": "高频任务集中在“比赛”和“我的球队”，球队首页保留年度数字与档案入口。",
        "screens": [("03-match-center", "比赛中心"), ("01-account-admin", "我的球队")],
        "points": [
            ("比赛入口", "查看近期比赛、历史比赛、排行榜，并进入赛事创建。"),
            ("球队入口", "汇总通知、名册、榜单、年鉴、个人档案和运营中心。"),
            ("球队首页", "展示下一场比赛、年度出场及进球助攻数据、最近赛事与核心球员。"),
        ],
    },
    {
        "eyebrow": "MEMBERSHIP & ACCESS",
        "title": "成员申请与四级权限",
        "subtitle": "访客、球员、管理员和队长具有清晰的查看及操作边界。",
        "screens": [("32-visitor-join", "访客申请"), ("33-account-player", "球员主页")],
        "points": [
            ("邀请加入", "访客填写真实姓名和邀请码，申请进入成员审核队列。"),
            ("成员功能", "审核通过后可查看队内信息、报名、通知、名册、年鉴和个人会费。"),
            ("管理权限", "管理员处理日常运营，队长拥有会费审批和管理员任命权限。"),
            ("个人设置", "成员可裁剪头像，并开启或关闭赛事订阅提醒。"),
        ],
    },
    {
        "eyebrow": "MATCH CENTER",
        "title": "近期比赛与历史档案",
        "subtitle": "同一入口覆盖赛前报名与赛后归档，赛事状态清晰可见。",
        "screens": [("03-match-center", "近期比赛"), ("04-match-history", "历史比赛")],
        "points": [
            ("近期比赛", "显示日期、时间、对手、地点、赛制、报名人数和报名状态。"),
            ("历史比赛", "沉淀比分、实际出场人数和赛果待补等状态。"),
            ("管理入口", "管理员可直接创建赛事，所有成员可进入排行榜。"),
        ],
    },
    {
        "eyebrow": "MATCH DETAIL",
        "title": "赛事详情、地图与队内信息",
        "subtitle": "比赛信息集中呈现，敏感字段按照成员身份开放。",
        "screens": [("05-match-detail", "赛事详情"), ("06-match-location", "地点与装备")],
        "points": [
            ("完整信息", "包含开球时间、集合时间、精确地点、赛制、球衣要求和赛事备注。"),
            ("地图能力", "支持保存地图坐标、展示位置卡片，并从小程序打开地图导航。"),
            ("信息保护", "精确地点、费用、内部备注和报名名单仅向审核通过的成员开放。"),
        ],
    },
    {
        "eyebrow": "TEAM REGISTRATION",
        "title": "队内报名与人员状态总览",
        "subtitle": "球员快速表达出席状态，管理层实时掌握阵容和特殊报名。",
        "screens": [("07-signup-options", "球员报名"), ("08-signup-manager", "管理名单")],
        "points": [
            ("三种状态", "支持参加、待定和缺席；待定状态需填写预计确认时间。"),
            ("报名类型", "支持本人参加、带人试训和带人凑脚，特殊报名进入审核。"),
            ("管理动作", "按状态筛选名单，审核特殊报名，提醒待定球员。"),
            ("群内协作", "一键复制微信群名单，报名备注可记录晚到、搭车等信息。"),
        ],
    },
    {
        "eyebrow": "TEAM SIGNAL",
        "title": "球队通知与定向发布",
        "subtitle": "通知留在站内形成记录，管理者可以按业务类型快速创建。",
        "screens": [("09-notices", "通知动态"), ("10-notice-publish", "发布通知")],
        "points": [
            ("通知模板", "覆盖赛事报名、报名动态、会费缴纳和待定确认。"),
            ("定向受众", "可发送给全体成员、未报名成员、待定球员、指定球员或管理层。"),
            ("关联信息", "赛事类通知可关联比赛，会费类通知可填写金额与截止日期。"),
            ("送达统计", "管理层可查看送达数量，成员点击后更新已读状态。"),
        ],
    },
    {
        "eyebrow": "SQUAD",
        "title": "现役阵容与历史成员",
        "subtitle": "名册按场上位置组织，同时保留历史成员与俱乐部荣誉档案。",
        "screens": [("11-player-roster", "现役球员"), ("12-alumni-roster", "历史成员")],
        "points": [
            ("位置分组", "前锋、中场、后卫、门将和待补资料分组展示。"),
            ("卡片信息", "展示号码、位置、赛季出场、进球和助攻。"),
            ("历史档案", "离队成员独立归档，支持董事长荣誉档案和长期投入记录。"),
        ],
    },
    {
        "eyebrow": "PLAYER PROFILE",
        "title": "球员档案、会费与荣誉",
        "subtitle": "个人数据、公开授权、会费信息和比赛记录集中在一个页面。",
        "screens": [("13-player-profile", "球员档案"), ("15-chairman-honor", "荣誉档案")],
        "points": [
            ("个人数据", "展示赛季出场、进球、助攻和最近相关赛事。"),
            ("隐私设置", "头像和比赛照片分别控制公开显示，成员状态清晰标注。"),
            ("个人会费", "本人及管理层可查看当前会费档次、生效期、原因与调整历史。"),
            ("荣誉记录", "历史成员可使用专属档案样式记录任职、投入和贡献。"),
        ],
    },
    {
        "eyebrow": "PLAYER MANAGEMENT & STATS",
        "title": "球员资料维护与数据榜单",
        "subtitle": "管理层维护球员基础信息，事实数据自动汇总为全队榜单。",
        "screens": [("16-player-edit", "编辑球员"), ("17-leaderboard", "排行榜")],
        "points": [
            ("资料维护", "可编辑姓名、号码、位置、效力年份、成员状态和公开授权。"),
            ("头像裁剪", "支持选择图片后拖动、缩放和框选裁剪区域。"),
            ("三类榜单", "出场榜、射手榜和助攻榜支持切换，并可查看全部记录。"),
            ("稳定排序", "同值时按较少出场及姓名规则稳定排序。"),
        ],
    },
    {
        "eyebrow": "THE YEARBOOK",
        "title": "球队年鉴与比赛章节",
        "subtitle": "历史赛事自动进入年鉴，形成可持续积累的球队数字档案。",
        "screens": [("18-yearbook", "年鉴列表"), ("19-yearbook-detail", "比赛章节")],
        "points": [
            ("赛季总览", "汇总比赛数和出场人次，精选章节使用重点样式。"),
            ("比赛章节", "记录日期、比分、出场、进球、助攻和比赛摘要。"),
            ("章节控制", "管理员可将章节设为精选、公开或内部。"),
        ],
    },
    {
        "eyebrow": "MATCH ALBUM",
        "title": "比赛相册与照片审核",
        "subtitle": "成员上传、管理审核、章节展示形成完整的照片沉淀流程。",
        "screens": [("20-yearbook-album", "相册入口"), ("27-photo-review", "照片审核")],
        "points": [
            ("成员上传", "每次最多选择 9 张照片，压缩后提交审核。"),
            ("完整相册", "管理员可保存百度网盘、夸克网盘或其他 HTTPS 相册链接。"),
            ("媒体审核", "支持通过、拒绝、设为封面、设为精选和隐藏。"),
            ("状态统计", "章节展示公开照片数量与等待审核数量。"),
        ],
    },
    {
        "eyebrow": "TEAM FINANCE",
        "title": "会费 Dashboard 与缴费台账",
        "subtitle": "按月查看应收、已收、待收和成员缴费状态。",
        "screens": [("21-fee-dashboard", "月度总览"), ("22-fee-ledger", "缴费台账")],
        "points": [
            ("月度看板", "显示应收、已收、待收、免缴人数、收缴率和已付未付人数。"),
            ("会费档次", "支持在莞人士、非在莞人士及高中生不同档次。"),
            ("台账操作", "按状态和档次筛选，快速标记已付或未付。"),
            ("自动规则", "学生寒暑假档次可按月份自动调整，个人页面同步显示。"),
        ],
    },
    {
        "eyebrow": "FEE APPROVAL",
        "title": "会费调整与队长审批",
        "subtitle": "管理员提交调整申请，队长确认后正式生效。",
        "screens": [("23-fee-adjustment", "管理员调整"), ("23b-fee-owner-review", "队长审核")],
        "points": [
            ("调整内容", "选择新档次、生效月份、结束月份，填写原因并决定是否通知球员。"),
            ("审批边界", "管理员调整进入待审队列，队长可确认通过或驳回。"),
            ("全程留痕", "申请人、申请时间、原因和审批结果进入操作记录。"),
        ],
    },
    {
        "eyebrow": "TEAM OPERATIONS",
        "title": "运营看板与管理员日志",
        "subtitle": "待办任务、快捷入口和敏感操作记录集中管理。",
        "screens": [("24-admin-dashboard", "运营看板"), ("25-admin-logs", "操作记录")],
        "points": [
            ("任务概览", "集中显示待审成员、会费审核、待审照片和赛事草稿。"),
            ("快捷操作", "直达会费、通知、赛事创建、赛后录入和榜单检查。"),
            ("操作日志", "记录动作、时间、摘要、备注和署名，仅管理层可见。"),
        ],
    },
    {
        "eyebrow": "REVIEW WORKFLOW",
        "title": "成员审核与媒体治理",
        "subtitle": "入队身份和公开照片均经过管理流程后生效。",
        "screens": [("26-member-review", "成员审核"), ("27-photo-review", "媒体审核")],
        "points": [
            ("档案绑定", "审核成员申请时绑定现有球员档案，避免重复身份。"),
            ("审核动作", "成员申请和照片均支持通过或拒绝。"),
            ("公开治理", "已通过照片仍可设置封面、精选或隐藏。"),
            ("迁移提醒", "历史数据待核对事项在运营中心持续提示。"),
        ],
    },
    {
        "eyebrow": "CAPTAIN CONTROL",
        "title": "管理员任命与权限治理",
        "subtitle": "队长独享最高权限，可维护球队管理人员。",
        "screens": [("28-admin-appointment", "任命入口"), ("29-admin-appointment-list", "成员名单")],
        "points": [
            ("队长专属", "管理员任命模块仅向队长显示。"),
            ("即时调整", "从成员名单任命或撤销管理员，权限状态同步更新。"),
            ("审计记录", "权限变更写入管理员操作记录，保留署名与时间。"),
        ],
    },
    {
        "eyebrow": "MATCH OPERATIONS",
        "title": "赛事创建与赛后统计",
        "subtitle": "赛前配置、状态管理和赛后事实数据发布形成业务闭环。",
        "screens": [("30-match-create", "创建赛事"), ("31-post-match", "赛后录入")],
        "points": [
            ("赛事配置", "维护标题、日期、时间、集合、赛制、地图地点、球衣、人数和备注。"),
            ("状态管理", "支持保存草稿、发布、截止报名和取消赛事。"),
            ("赛后事实", "录入比分、实际出场、进球和助攻；比分可暂时留空。"),
            ("数据校验", "助攻数高于进球数时阻止发布，发布后重算球员与榜单数据。"),
        ],
    },
]


def ensure_images() -> None:
    OPTIMIZED.mkdir(parents=True, exist_ok=True)
    required = {name for section in SECTIONS for name, _ in section["screens"]}
    required.update({"02-team-home", "03-match-center", "21-fee-dashboard", "24-admin-dashboard"})
    missing = [name for name in sorted(required) if not (SCREENS / f"{name}.png").exists()]
    if missing:
        raise FileNotFoundError(f"Missing screenshots: {', '.join(missing)}")
    for name in required:
        source = SCREENS / f"{name}.png"
        target = OPTIMIZED / f"{name}.jpg"
        with Image.open(source) as image:
            image = image.convert("RGB")
            image.thumbnail((430, 932), Image.Resampling.LANCZOS)
            image.save(target, "JPEG", quality=63, optimize=True, progressive=True)


def register_pdf_fonts() -> None:
    pdfmetrics.registerFont(TTFont("CN", r"C:\Windows\Fonts\msyh.ttc"))
    pdfmetrics.registerFont(TTFont("CN-Bold", r"C:\Windows\Fonts\msyhbd.ttc"))


def wrap_text(text: str, font: str, size: float, width: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for character in text:
        if character == "\n":
            lines.append(current)
            current = ""
        elif pdfmetrics.stringWidth(current + character, font, size) <= width:
            current += character
        else:
            lines.append(current)
            current = character
    if current:
        lines.append(current)
    return lines


def pdf_paragraph(c: canvas.Canvas, text: str, x: float, y: float, width: float,
                  size: float = 8.4, color=MUTED, leading: float = 12,
                  font: str = "CN") -> float:
    c.setFont(font, size)
    c.setFillColor(color)
    for line in wrap_text(text, font, size, width):
        c.drawString(x, y, line)
        y -= leading
    return y


def gradient(c: canvas.Canvas, top: str = "#06133D", bottom: str = "#2D6DDB") -> None:
    start, end = HexColor(top), HexColor(bottom)
    for index in range(100):
        ratio = index / 99
        color = Color(
            start.red + (end.red - start.red) * ratio,
            start.green + (end.green - start.green) * ratio,
            start.blue + (end.blue - start.blue) * ratio,
        )
        c.setFillColor(color)
        c.rect(0, index * PH / 100, PW, PH / 100 + 1, fill=1, stroke=0)
    c.setFillColor(Color(0.55, 0.75, 1, alpha=0.10))
    c.circle(PW - 55, PH - 30, 180, fill=1, stroke=0)


def rounded(c: canvas.Canvas, x: float, y: float, width: float, height: float,
            fill=white, stroke=LINE, radius: float = 12, shadow: bool = True) -> None:
    if shadow:
        c.setFillColor(Color(0.02, 0.08, 0.25, alpha=0.09))
        c.roundRect(x + 3, y - 4, width, height, radius, fill=1, stroke=0)
    c.setFillColor(fill)
    c.setStrokeColor(stroke)
    c.setLineWidth(0.8)
    c.roundRect(x, y, width, height, radius, fill=1, stroke=1)


def pdf_base(c: canvas.Canvas, page_number: int, section: str, dark: bool = False) -> None:
    if dark:
        gradient(c)
    else:
        c.setFillColor(PAPER)
        c.rect(0, 0, PW, PH, fill=1, stroke=0)
        c.setFillColor(Color(0.12, 0.35, 0.85, alpha=0.055))
        c.circle(PW - 25, PH - 25, 165, fill=1, stroke=0)
    c.setFont("CN-Bold", 8)
    c.setFillColor(ICE if dark else ROYAL)
    c.drawString(36, PH - 25, "大湾区流浪者  /  GBA RANGERS")
    c.setFont("CN", 7.5)
    c.setFillColor(Color(1, 1, 1, alpha=0.66) if dark else MUTED)
    c.drawRightString(PW - 36, PH - 25, section)
    c.setStrokeColor(Color(1, 1, 1, alpha=0.18) if dark else LINE)
    c.line(36, 24, PW - 36, 24)
    c.setFont("CN", 7)
    c.setFillColor(Color(1, 1, 1, alpha=0.58) if dark else MUTED)
    c.drawString(36, 12, "产品手册  ·  当前项目真实运行界面  ·  2026.07.29")
    c.drawRightString(PW - 36, 12, f"{page_number:02d}")


def pdf_title(c: canvas.Canvas, eyebrow: str, title: str, subtitle: str,
              dark: bool = False) -> None:
    c.setFillColor(ICE if dark else ROYAL)
    c.setFont("CN-Bold", 8.5)
    c.drawString(48, PH - 62, eyebrow)
    c.setFillColor(white if dark else INK)
    c.setFont("CN-Bold", 22)
    c.drawString(48, PH - 92, title)
    c.setFillColor(Color(1, 1, 1, alpha=0.72) if dark else MUTED)
    c.setFont("CN", 9.3)
    c.drawString(48, PH - 113, subtitle)


def draw_phone(c: canvas.Canvas, name: str, x: float, y: float,
               max_width: float, max_height: float, caption: str = "") -> None:
    source = OPTIMIZED / f"{name}.jpg"
    with Image.open(source) as image:
        width, height = image.size
    scale = min((max_width - 10) / width, (max_height - 10) / height)
    draw_width, draw_height = width * scale, height * scale
    frame_width, frame_height = draw_width + 10, draw_height + 10
    frame_x = x + (max_width - frame_width) / 2
    frame_y = y + (max_height - frame_height) / 2
    rounded(c, frame_x, frame_y, frame_width, frame_height,
            fill=HexColor("#101827"), stroke=Color(1, 1, 1, alpha=0.42), radius=14)
    c.drawImage(ImageReader(str(source)), frame_x + 5, frame_y + 5,
                draw_width, draw_height, preserveAspectRatio=True)
    if caption:
        c.setFillColor(INK)
        c.setFont("CN-Bold", 8)
        c.drawCentredString(x + max_width / 2, y - 10, caption)


def pdf_feature_card(c: canvas.Canvas, x: float, y: float, width: float,
                     number: str, title: str, body: str, accent=ROYAL) -> None:
    rounded(c, x, y, width, 67)
    c.setFillColor(accent)
    c.circle(x + 24, y + 43, 11, fill=1, stroke=0)
    c.setFillColor(white)
    c.setFont("CN-Bold", 7.5)
    c.drawCentredString(x + 24, y + 40.5, number)
    c.setFillColor(INK)
    c.setFont("CN-Bold", 10.5)
    c.drawString(x + 44, y + 44, title)
    pdf_paragraph(c, body, x + 16, y + 23, width - 32, 8.1, MUTED, 11.5)


def build_pdf() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    register_pdf_fonts()
    c = canvas.Canvas(str(PDF_OUT), pagesize=(PW, PH), pageCompression=1)
    c.setTitle("大湾区流浪者小程序产品手册 - 2026 年 7 月")
    c.setAuthor("大湾区流浪者")
    c.setSubject("球队赛事、成员、年鉴、会费与运营管理功能手册")

    gradient(c, "#04113F", "#174FC7")
    c.setFillColor(Color(1, 1, 1, alpha=0.075))
    c.roundRect(42, 38, PW - 84, PH - 76, 28, fill=1, stroke=0)
    c.setFillColor(ICE)
    c.setFont("CN-Bold", 10)
    c.drawString(64, PH - 82, "GBA RANGERS  /  PRODUCT MANUAL")
    c.setFillColor(white)
    c.setFont("CN-Bold", 31)
    c.drawString(64, PH - 140, "大湾区流浪者小程序")
    c.setFont("CN-Bold", 23)
    c.drawString(64, PH - 178, "最新产品手册")
    c.setFillColor(Color(1, 1, 1, alpha=0.74))
    c.setFont("CN", 10.5)
    c.drawString(64, PH - 210, "赛事 · 成员 · 年鉴 · 会费 · 球队运营")
    c.setFont("CN", 8)
    c.drawString(64, 62, "基于当前项目成果整理，界面截图来自 2026 年 7 月 29 日本地演示构建。")
    if LOGO.exists():
        c.saveState()
        c.setFillAlpha(0.13)
        c.drawImage(str(LOGO), 48, 78, 230, 230, mask="auto", preserveAspectRatio=True)
        c.restoreState()
    draw_phone(c, "03-match-center", 418, 76, 118, 348)
    draw_phone(c, "21-fee-dashboard", 542, 58, 118, 348)
    draw_phone(c, "24-admin-dashboard", 666, 76, 118, 348)
    c.showPage()

    pdf_base(c, 2, "产品全景", dark=False)
    pdf_title(c, "PRODUCT MAP", "一套围绕球队日常运营的数字工具",
              "成员从报名和通知开始使用，管理层在同一体系内完成赛事、数据、会费与内容治理。")
    x_positions = [48, 238, 428, 618]
    y_positions = [378, 278, 178, 78]
    accents = [ROYAL, BLUE, NAVY, GOLD]
    for index, (number, title, body) in enumerate(FEATURE_INDEX):
        column = index % 4
        row = index // 4
        pdf_feature_card(c, x_positions[column], y_positions[row], 172, number, title, body, accents[column])
    c.showPage()

    page_number = 3
    for section in SECTIONS:
        pdf_base(c, page_number, section["title"], dark=False)
        pdf_title(c, section["eyebrow"], section["title"], section["subtitle"])
        first, second = section["screens"]
        draw_phone(c, first[0], 48, 70, 176, 374, first[1])
        draw_phone(c, second[0], 228, 70, 176, 374, second[1])
        c.setFillColor(INK)
        c.setFont("CN-Bold", 13.5)
        c.drawString(438, 438, "功能说明")
        card_y = 352
        accents = [ROYAL, BLUE, GREEN, GOLD]
        for index, (title, body) in enumerate(section["points"], 1):
            pdf_feature_card(c, 438, card_y, 350, f"{index:02d}", title, body, accents[(index - 1) % len(accents)])
            card_y -= 80
        c.showPage()
        page_number += 1

    pdf_base(c, page_number, "交付边界", dark=True)
    pdf_title(c, "DELIVERY & SECURITY", "当前交付状态与数据安全边界",
              "本地演示覆盖全部核心流程，生产运行依赖 CloudBase 环境和云函数部署。", dark=True)
    cards = [
        ("客户端", "Taro 4、React、TypeScript；微信小程序与 H5 共享业务代码和视觉体系。"),
        ("数据模式", "本地模式使用完整演示数据；生产模式已配置 CloudBase 环境。"),
        ("权限边界", "公开 DTO 不包含精确地点、费用、报名名单、内部备注和未审核媒体。"),
        ("事实数据", "单场出场与球员比赛统计是事实源，累计数据通过重算生成。"),
        ("审计能力", "赛事、成员、媒体、会费和权限变更均保留操作记录。"),
        ("上线准备", "生产启用前需创建集合与索引、部署七个云函数，并完成真机验收。"),
    ]
    for index, (title, body) in enumerate(cards):
        column = index % 2
        row = index // 2
        x = 56 + column * 378
        y = 350 - row * 112
        rounded(c, x, y, 350, 88, fill=Color(1, 1, 1, alpha=0.11),
                stroke=Color(1, 1, 1, alpha=0.22), shadow=False)
        c.setFillColor(ICE)
        c.setFont("CN-Bold", 8)
        c.drawString(x + 18, y + 64, f"0{index + 1}")
        c.setFillColor(white)
        c.setFont("CN-Bold", 12)
        c.drawString(x + 48, y + 62, title)
        pdf_paragraph(c, body, x + 18, y + 39, 314, 8.2, Color(1, 1, 1, alpha=0.72), 12)
    c.setFillColor(Color(1, 1, 1, alpha=0.70))
    c.setFont("CN", 8.5)
    c.drawString(56, 54, "建议上线前使用微信开发者工具官方模拟器验收，并在 iPhone 与 Android 各完成一次真机测试。")
    c.showPage()

    c.save()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr()
    fonts = run._element.rPr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.rPr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_tokens(paragraph, before: float = 0, after: float = 6,
                         line: float = 1.25, keep_with_next: bool = False) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_with_next = keep_with_next


def add_docx_text(document: Document, text: str, size: float = 11,
                  bold: bool = False, color: str = "0A173A",
                  alignment=WD_ALIGN_PARAGRAPH.LEFT, before: float = 0,
                  after: float = 6) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = alignment
    set_paragraph_tokens(paragraph, before, after, 1.25)
    run = paragraph.add_run(text)
    set_run_font(run, size, bold, color)


def add_docx_image_row(document: Document, screens: list[tuple[str, str]]) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(5)
    for index, (name, _) in enumerate(screens):
        if index:
            spacer = paragraph.add_run("     ")
            set_run_font(spacer, 3)
        run = paragraph.add_run()
        run.add_picture(str(OPTIMIZED / f"{name}.jpg"), width=Inches(2.42))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(8)
    for index, (_, label) in enumerate(screens):
        if index:
            spacer = caption.add_run("                                      ")
            set_run_font(spacer, 8)
        run = caption.add_run(label)
        set_run_font(run, 8.5, True, "5F6D89")


def add_docx_bullets(document: Document, points: list[tuple[str, str]]) -> None:
    for title, body in points:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.38)
        paragraph.paragraph_format.first_line_indent = Inches(-0.19)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        paragraph.paragraph_format.line_spacing = 1.25
        title_run = paragraph.add_run(f"{title}：")
        set_run_font(title_run, 10.5, True, "123BB4")
        body_run = paragraph.add_run(body)
        set_run_font(body_run, 10.5, False, "0A173A")


def add_page_break(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def build_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Title", 30, "0A173A", 0, 8),
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Title"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Eyebrow" not in styles:
        eyebrow_style = styles.add_style("Eyebrow", WD_STYLE_TYPE.PARAGRAPH)
    else:
        eyebrow_style = styles["Eyebrow"]
    eyebrow_style.font.name = "Calibri"
    eyebrow_style.font.size = Pt(8.5)
    eyebrow_style.font.bold = True
    eyebrow_style.font.color.rgb = RGBColor.from_string("123BB4")
    eyebrow_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    eyebrow_style.paragraph_format.space_after = Pt(4)
    eyebrow_style.paragraph_format.keep_with_next = True

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    header_run = header_p.add_run("GBA RANGERS  /  大湾区流浪者")
    set_run_font(header_run, 8, True, "63708D")
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer_p.add_run("产品手册  ·  2026.07.29")
    set_run_font(footer_run, 8, False, "63708D")

    add_docx_text(document, "GBA RANGERS  /  PRODUCT MANUAL", 9, True, "123BB4", before=52, after=16)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("大湾区流浪者小程序")
    set_run_font(title_run, 30, False, "0A173A")
    add_docx_text(document, "最新产品手册", 22, True, "123BB4", after=12)
    add_docx_text(document, "赛事 · 成员 · 年鉴 · 会费 · 球队运营", 12, False, "5F6D89", after=26)
    cover_screens = [
        ("03-match-center", "比赛中心"),
        ("21-fee-dashboard", "会费管理"),
        ("24-admin-dashboard", "运营中心"),
    ]
    cover_p = document.add_paragraph()
    cover_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, (name, _) in enumerate(cover_screens):
        if index:
            cover_p.add_run("   ")
        cover_p.add_run().add_picture(str(OPTIMIZED / f"{name}.jpg"), width=Inches(1.55))
    add_docx_text(document, "基于当前项目成果整理，截图来自 2026 年 7 月 29 日本地演示构建。", 9, False, "63708D",
                  alignment=WD_ALIGN_PARAGRAPH.CENTER, before=14, after=0)
    add_page_break(document)

    document.add_paragraph("产品功能索引", style="Heading 1")
    add_docx_text(document, "当前版本覆盖 16 个核心模块，成员端与管理端共享同一套数据和权限体系。", 11, False, "5F6D89", after=12)
    for number, title, body in FEATURE_INDEX:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.left_indent = Inches(0.42)
        paragraph.paragraph_format.first_line_indent = Inches(-0.21)
        paragraph.paragraph_format.space_after = Pt(7)
        paragraph.paragraph_format.line_spacing = 1.25
        run = paragraph.add_run(f"{title}：")
        set_run_font(run, 11, True, "123BB4")
        run = paragraph.add_run(body)
        set_run_font(run, 11, False, "0A173A")
    add_page_break(document)

    for section_data in SECTIONS:
        eyebrow = document.add_paragraph(style="Eyebrow")
        eyebrow.add_run(section_data["eyebrow"])
        heading = document.add_paragraph(section_data["title"], style="Heading 1")
        add_docx_text(document, section_data["subtitle"], 10.5, False, "5F6D89", after=8)
        add_docx_image_row(document, section_data["screens"])
        add_docx_bullets(document, section_data["points"])
        add_page_break(document)

    eyebrow = document.add_paragraph(style="Eyebrow")
    eyebrow.add_run("DELIVERY & SECURITY")
    document.add_paragraph("当前交付状态与数据安全边界", style="Heading 1")
    add_docx_text(document, "本地演示覆盖全部核心流程，生产运行依赖 CloudBase 环境和云函数部署。", 10.5, False, "5F6D89", after=14)
    delivery_points = [
        ("客户端", "Taro 4、React、TypeScript；微信小程序与 H5 共享业务代码和视觉体系。"),
        ("数据模式", "本地模式使用完整演示数据；生产模式已配置 CloudBase 环境。"),
        ("权限边界", "公开 DTO 不包含精确地点、费用、报名名单、内部备注和未审核媒体。"),
        ("事实数据", "单场出场与球员比赛统计是事实源，累计数据通过重算生成。"),
        ("审计能力", "赛事、成员、媒体、会费和权限变更均保留操作记录。"),
        ("上线准备", "生产启用前需创建集合与索引、部署七个云函数，并完成真机验收。"),
    ]
    add_docx_bullets(document, delivery_points)
    add_docx_text(document, "上线验收建议", 13, True, "2E74B5", before=16, after=8)
    add_docx_bullets(document, [
        ("官方模拟器", "确认页面尺寸、字体、底部导航、返回路径、输入组件和角色权限。"),
        ("真机预览", "至少覆盖一台 iPhone 和一台 Android，检查安全区、授权、相册、相机与触控。"),
        ("云端联调", "部署数据库与云函数后，验证登录身份、权限规则、媒体上传和订阅消息。"),
    ])

    if document.paragraphs and not document.paragraphs[-1].text:
        last = document.paragraphs[-1]._element
        last.getparent().remove(last)

    document.core_properties.title = "大湾区流浪者小程序产品手册 - 2026 年 7 月"
    document.core_properties.subject = "球队赛事、成员、年鉴、会费与运营管理功能手册"
    document.core_properties.author = "大湾区流浪者"
    document.core_properties.keywords = "GBA RANGERS, 微信小程序, 产品手册"
    OUTPUT.mkdir(parents=True, exist_ok=True)
    document.save(DOCX_OUT)


def main() -> None:
    ensure_images()
    build_pdf()
    build_docx()
    print(PDF_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    main()
